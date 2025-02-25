from fpdf import FPDF
import os
import docx

class PDF(FPDF):
    def chapter_title(self, title):
        self.set_font('DejaVu', 'B', 14)
        self.cell(0, 10, title, 0, 1, 'L')
        self.ln(4)  # Add space after the title

    def chapter_body(self, body):
        self.set_font('DejaVu', '', 12)
        self.multi_cell(0, 10, body)
        self.ln()  # Add space after the body

def create_resume_template(output_path, template_path, optimized_resume, contact_information, relevant_experiences, work_experiences, education_section, certifications_and_licenses):
    pdf = PDF()
    pdf.add_page()
    pdf.add_font('DejaVu', '', os.path.join(os.path.dirname(__file__), 'dejavu-fonts-ttf-2.37/ttf/DejaVuSansCondensed.ttf'), uni=True)
    pdf.set_font('DejaVu', '', 12)

    sections = {
        'Contact Information': contact_information,
        'Professional Summary': '',
        'Relevant Experience': relevant_experiences,
        'Work Experience': work_experiences,
        'Education': education_section,
        'Skills': '',
        'Certifications and Licenses': certifications_and_licenses
    }

    # Extract sections from optimized_resume
    current_section = None
    for line in optimized_resume.split('\n'):
        line = line.strip()
        if line in sections:
            current_section = line
        elif current_section:
            sections[current_section] += line + '\n'

    # Add sections to the PDF
    for title, body in sections.items():
        if body.strip():
            pdf.chapter_title(title)  # Bold section title
            pdf.chapter_body(body)

    pdf.output(output_path, 'F')

def create_resume_docx(resume_content, output_path):
    doc = docx.Document()
    doc.add_paragraph(resume_content)
    doc.save(output_path)
